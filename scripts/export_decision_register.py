"""Export the checked workbook as a short index plus lossless readable evidence."""

import argparse
from datetime import date, datetime
from pathlib import Path

from decision_register import DETAIL, TEAM, check_book
from openpyxl import load_workbook


def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def inline(value):
    return text(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def export(workbook, output):
    output = Path(output)
    details = output.with_name("DETAILS.md")
    if output.suffix.lower() != ".md":
        raise ValueError("Summary output must be a Markdown file")
    if output.resolve() == details.resolve():
        raise ValueError("Summary output cannot be named DETAILS.md")
    if Path(workbook).resolve() in {output.resolve(), details.resolve()}:
        raise ValueError("Generated copies must not overwrite the workbook")
    book = load_workbook(workbook, data_only=False)
    try:
        check_book(book)
        team = book[TEAM]
        source = book[DETAIL]
        headings = [cell.value for cell in source[4][:22]]
        summaries = [
            list(row)
            for row in team.iter_rows(min_row=5, max_col=6, values_only=True)
            if row[0] is not None
        ]
        records = [
            list(row)
            for row in source.iter_rows(min_row=5, max_col=22, values_only=True)
            if row[0] is not None
        ]
        stamp = datetime.now().astimezone().isoformat(timespec="seconds")
        intro = (
            f"Generated from `{workbook}` at {stamp}. Edit the workbook, not this copy."
        )
        lines = [
            "# Decision register - team view",
            "",
            intro,
            "",
            f"{len(summaries)} decisions. Full source fields and history: [DETAILS.md](DETAILS.md).",
            "",
            "| ID | Topic | Current position | Status | Owner | Next / due |",
            "|---|---|---|---|---|---|",
        ]
        for row in summaries:
            key = str(row[0])
            row[0] = f"[{key}](DETAILS.md#{key.lower()})"
            lines.append("| " + " | ".join(inline(value) for value in row) + " |")
        evidence = [
            "# Decision register - full evidence",
            "",
            intro,
            "",
            "All populated source fields, without clipping or latest-entry extraction. Earlier entries may be superseded; read chronology before using them. Formula expressions are preserved, not evaluated.",
            "",
        ]
        for row in records:
            evidence += [f"## {text(row[0])}", "", f"**{text(row[1])}**", ""]
            for heading, value in zip(headings[2:], row[2:]):
                if value is not None:
                    evidence += [f"### {heading}", "", text(value), ""]
        # Prepare both complete strings before publishing either copy.
        summary_text = "\n".join(lines) + "\n"
        detail_text = "\n".join(evidence) + "\n"
        output.parent.mkdir(parents=True, exist_ok=True)
        details.write_text(detail_text, encoding="utf-8", newline="\n")
        output.write_text(summary_text, encoding="utf-8", newline="\n")
        return len(summaries)
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    count = export(args.workbook, args.output)
    print(f"{count} decisions -> {args.output} + DETAILS.md (unclipped)")


if __name__ == "__main__":
    main()
