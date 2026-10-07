"""Build, check and export a concise view of a canonical decision register.

Run with uv run --with openpyxl scripts/decision_register.py --help.
Build defaults to a sibling <stem>.team.xlsx; source files are never promoted.
"""

import argparse
import json
import os
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

DETAIL = "Decision Register"
TEAM = "Team View"
DETAIL_HEADERS = [
    "Decision ID",
    "Title",
    "Description / Context",
    "Decision Made",
    "Rationale / Basis",
    "Status",
    "Priority",
    "Phase",
    "Discipline / Category",
    "Owner",
    "Decision Authority",
    "Stakeholders / Affected",
    "Date Raised",
    "Needed-By Date",
    "Date Decided",
    "Cost Impact",
    "Schedule Impact",
    "Dependencies",
    "Superseded By",
    "Open Questions",
    "Reference / Link",
    "Comments / Notes",
]
HEADERS = ["ID", "Topic", "Current position", "Status", "Owner", "Next / due"]
WIDTHS = (8, 25, 49, 13, 14, 37)
TITLE = "Decision register - current summaries"
HELP = (
    "Current position, not history. Filter Status; click an ID for canonical details."
)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def identifier(value):
    require(
        isinstance(value, (str, int)) and not isinstance(value, bool),
        "ID must be text or an integer",
    )
    result = str(value)
    require(
        bool(result.strip()) and result == result.strip(),
        "ID must be nonempty with no outer whitespace",
    )
    return result


def text_value(value, label, limit=None, empty=False):
    require(isinstance(value, str), f"{label} must be a string")
    require(empty or bool(value.strip()), f"{label} must be nonempty")
    require(limit is None or len(value) <= limit, f"{label} exceeds {limit} characters")
    require(not any(c in value for c in "\r\n\t"), f"{label} must be one paragraph")


def canonical(book):
    require(DETAIL in book.sheetnames, f"Missing {DETAIL}")
    sheet = book[DETAIL]
    require(
        [sheet.cell(4, c).value for c in range(1, 23)] == DETAIL_HEADERS,
        "Unexpected canonical headers at A4:V4",
    )
    records = {}
    for cells in sheet.iter_rows(min_row=5, max_col=22):
        if all(cell.value is None for cell in cells):
            continue
        require(
            all(cells[c].data_type != "f" for c in (0, 5, 9)),
            "Canonical ID, Status and Owner must be literal values",
        )
        key = identifier(cells[0].value)
        require(key not in records, f"Duplicate canonical ID: {key}")
        status, owner = cells[5].value, cells[9].value or ""
        text_value(status, f"{key} status")
        require(
            status in {"Open", "In Review", "Decided", "Deferred", "Superseded"},
            f"{key} unknown status: {status}",
        )
        text_value(owner, f"{key} owner", 40, empty=True)
        records[key] = (cells[0].row, cells[0].value, status, owner)
    require(bool(records), "Canonical register has no records")
    return records


def validate_summaries(items, records):
    require(isinstance(items, list), "Summaries must be a JSON list")
    mapped = {}
    for item in items:
        require(
            isinstance(item, dict) and set(item) == {"id", "topic", "position", "next"},
            "Each summary needs exactly id, topic, position, next",
        )
        key = identifier(item["id"])
        require(key not in mapped, f"Duplicate summary ID: {key}")
        require(key in records, f"Unknown summary ID: {key}")
        text_value(item["topic"], f"{key} topic", 55)
        text_value(item["position"], f"{key} position", 150)
        text_value(
            item["next"],
            f"{key} next",
            100,
            records[key][2].strip().casefold() in {"decided", "superseded"},
        )
        mapped[key] = item
    require(
        set(mapped) == set(records),
        f"Missing summary IDs: {sorted(set(records) - set(mapped))}",
    )
    return mapped


def team_layout(sheet):
    require(not sheet.merged_cells.ranges, "Team View must not contain merges")
    require(
        [sheet.cell(4, c).value for c in range(1, 7)] == HEADERS,
        "Unexpected Team View headers",
    )
    for row in sheet.iter_rows():
        for cell in row:
            require(cell.data_type != "f", f"Team View formula at {cell.coordinate}")
            require(
                cell.value is None or cell.column <= 6,
                "Unexpected content outside Team View columns",
            )


def extract(book, records, strict=True):
    require(TEAM in book.sheetnames, "Missing Team View")
    sheet = book[TEAM]
    team_layout(sheet)
    items = []
    for row in sheet.iter_rows(min_row=5, max_col=6):
        if all(cell.value is None for cell in row):
            continue
        key = identifier(row[0].value)
        require(key in records, f"Unknown Team View ID: {key}")
        if strict:
            require(
                row[3].value == records[key][2], f"{key} status differs from canonical"
            )
            require(
                (row[4].value or "") == records[key][3],
                f"{key} owner differs from canonical",
            )
        items.append(
            {
                "id": row[0].value,
                "topic": row[1].value,
                "position": row[2].value,
                "next": row[5].value if row[5].value is not None else "",
            }
        )
    require(
        len({identifier(item["id"]) for item in items}) == len(items),
        "Duplicate Team View IDs",
    )
    if strict:
        validate_summaries(items, records)
    return items


def check_book(book):
    records = canonical(book)
    items = extract(book, records)
    sheet = book[TEAM]
    require(
        book.active == sheet and book.worksheets[0] == sheet,
        "Team View must be first and active",
    )
    require(sheet.sheet_state == "visible", "Team View must be visible")
    require((sheet.sheet_view.zoomScale or 100) >= 85, "Zoom must be at least 85%")
    require(sheet.freeze_panes in {"A5", "B5"}, "Freeze panes must be A5 or B5")
    require(not sheet.sheet_format.zeroHeight, "Default rows must not be hidden")
    require(
        sheet.sheet_view.showOutlineSymbols is False,
        "Grouping indicators must be cleared",
    )
    for dimension in list(sheet.row_dimensions.values()) + list(
        sheet.column_dimensions.values()
    ):
        require(not dimension.hidden, "Hidden rows or columns")
        require(
            not dimension.outlineLevel and not dimension.collapsed,
            "Grouping must be cleared",
        )
    for col, width in zip("ABCDEF", WIDTHS):
        require(
            sheet.column_dimensions[col].width == width,
            f"Column {col} must have width {width}",
        )
    end = 4 + len(items)
    require(sheet.auto_filter.ref == f"A4:F{end}", "Autofilter must cover every record")
    require(
        not sheet.auto_filter.filterColumn and sheet.auto_filter.sortState is None,
        "Stale filter or sort criteria",
    )
    require(not sheet.sheet_properties.filterMode, "Stale filter mode")
    require(not sheet.tables, "Unexpected Team View tables")
    for row in range(4, end + 1):
        height = sheet.row_dimensions[row].height
        if height is None:
            height = sheet.sheet_format.defaultRowHeight
        require(
            height is not None and 36 <= height <= 48, f"Row {row} height must be 36-48"
        )
        for cell in sheet[row][:6]:
            require(
                cell.alignment.wrap_text and cell.alignment.vertical == "top",
                f"{cell.coordinate} must wrap and align top",
            )
            require(
                cell.font.sz == 11 and cell.font.name == "Calibri",
                f"{cell.coordinate} font must be Calibri 11pt",
            )
            require(
                not cell.alignment.shrink_to_fit,
                f"{cell.coordinate} must not shrink to fit",
            )
        if row >= 5:
            key = identifier(sheet.cell(row, 1).value)
            link = sheet.cell(row, 1).hyperlink
            require(
                link is not None
                and link.target is None
                and link.location == f"'{DETAIL}'!A{records[key][0]}",
                f"{key} needs an internal detail link",
            )
    require(
        not any(
            cell.value is not None
            for row in sheet.iter_rows(min_row=end + 1)
            for cell in row
        ),
        "Unexpected trailing records",
    )
    return items


def distinct(source, destination):
    require(
        source.resolve() != destination.resolve(),
        "Output must differ from input workbook",
    )
    require(
        not destination.exists() or not os.path.samefile(source, destination),
        "Output must differ from input workbook",
    )


def build(workbook, summaries, output=None):
    source = Path(workbook)
    destination = (
        Path(output) if output else source.with_name(source.stem + ".team.xlsx")
    )
    distinct(source, destination)
    require(
        source.suffix.lower() == destination.suffix.lower() == ".xlsx",
        "Build requires .xlsx input and output",
    )
    require(
        Path(summaries).resolve() != destination.resolve(),
        "Output must not overwrite summaries",
    )
    book = load_workbook(source)
    temporary = None
    try:
        records = canonical(book)
        with open(summaries, encoding="utf-8") as stream:
            mapped = validate_summaries(json.load(stream), records)
        comments = {}
        header_comments = {}
        if TEAM in book.sheetnames:
            existing = book[TEAM]
            team_layout(existing)
            require(
                existing["A1"].value == TITLE and existing["A2"].value == HELP,
                "Unrecognized existing Team View",
            )
            require(
                all(
                    cell.value is None
                    for row in existing.iter_rows(min_row=1, max_row=3)
                    for cell in row
                    if cell.coordinate not in {"A1", "A2"}
                ),
                "Unexpected Team View front matter",
            )
            extract(book, records, strict=False)
            require(
                not existing.tables and not existing._charts and not existing._images,
                "Team View has custom objects; preserve them before rebuilding",
            )
            for row in existing:
                for cell in row:
                    if cell.comment is None:
                        continue
                    require(
                        cell.column <= 6,
                        "Comment outside Team View columns; relocate it before rebuilding",
                    )
                    if cell.row < 5:
                        header_comments[cell.coordinate] = copy(cell.comment)
                    else:
                        key = identifier(existing.cell(cell.row, 1).value)
                        require(
                            key in records,
                            "Comment on an unkeyed row; relocate it before rebuilding",
                        )
                        comments[key, cell.column] = copy(cell.comment)
            del book[TEAM]
        sheet = book.create_sheet(TEAM, 0)
        sheet["A1"] = TITLE
        sheet["A2"] = HELP
        sheet["A1"].font = Font(size=16, bold=True, color="17365D")
        sheet["A2"].font = Font(size=11, color="595959")
        sheet.row_dimensions[1].height = 24
        sheet.row_dimensions[2].height = 20
        for col, width in zip("ABCDEF", WIDTHS):
            sheet.column_dimensions[col].width = width
        for col, value in enumerate(HEADERS, 1):
            sheet.cell(4, col, value)
        for row, (key, (detail_row, raw_id, status, owner)) in enumerate(
            records.items(), 5
        ):
            item = mapped[key]
            for col, value in enumerate(
                (raw_id, item["topic"], item["position"], status, owner, item["next"]),
                1,
            ):
                cell = sheet.cell(row, col, value)
                if (key, col) in comments:
                    cell.comment = comments[key, col]
                if isinstance(value, str):
                    cell.data_type = (
                        "s"  # Summaries are literal text, even when starting with '='.
                    )
            sheet.cell(row, 1).hyperlink = Hyperlink(
                ref=f"A{row}", location=f"'{DETAIL}'!A{detail_row}"
            )
        for row in sheet.iter_rows(min_row=4, max_row=4 + len(records), max_col=6):
            sheet.row_dimensions[row[0].row].height = 48
            for cell in row:
                header = cell.row == 4
                cell.font = Font(
                    name="Calibri",
                    size=11,
                    bold=header,
                    color="FFFFFF" if header else "17365D",
                )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill(
                    "solid",
                    fgColor="17365D"
                    if header
                    else ("EDF2F7" if cell.row % 2 else "FFFFFF"),
                )
        for coordinate, comment in header_comments.items():
            sheet[coordinate].comment = comment
        sheet.auto_filter.ref = f"A4:F{4 + len(records)}"
        sheet.print_area = f"A1:F{4 + len(records)}"
        sheet.print_title_rows = "1:4"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.freeze_panes = "B5"
        sheet.sheet_view.zoomScale = 85
        sheet.sheet_view.showOutlineSymbols = False
        book.active = 0
        check_book(book)
        with tempfile.NamedTemporaryFile(
            dir=destination.parent, suffix=".xlsx", delete=False
        ) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        check(temporary)  # Validate actual serialized values before publishing output.
        os.replace(temporary, destination)
        return destination
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def check(workbook):
    book = load_workbook(workbook)
    try:
        return check_book(book)
    finally:
        book.close()


def check_rendering(workbook):
    """Measure wrapped rows with a spreadsheet engine, without saving its copy.

    Optional dependency: aspose-cells-python. Evaluation rendering can watermark
    images; this check publishes neither an image nor an engine-written workbook.
    """
    from aspose.cells import Workbook

    check(workbook)
    book = Workbook(str(Path(workbook).resolve()))
    sheet = book.worksheets.get(TEAM)
    last = sheet.cells.max_data_row
    heights = {row: sheet.cells.get_row_height(row) for row in range(4, last + 1)}
    sheet.auto_fit_rows()
    overflow = [
        f"{sheet.cells.get(row, 0).string_value}: needs {sheet.cells.get_row_height(row):g}pt, has {height:g}pt"
        for row, height in heights.items()
        if sheet.cells.get_row_height(row) > height
    ]
    require(not overflow, "Wrapped text would clip: " + "; ".join(overflow))
    return len(heights)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    builder = commands.add_parser("build")
    builder.add_argument("workbook", type=Path)
    builder.add_argument("--summaries", type=Path, required=True)
    builder.add_argument("--output", type=Path)
    checker = commands.add_parser("check")
    checker.add_argument("workbook", type=Path)
    checker.add_argument(
        "--render",
        action="store_true",
        help="Also measure wrapping with aspose-cells-python; never saves the engine copy",
    )
    exporter = commands.add_parser(
        "export-summaries", help="Export validated current summaries; default stdout"
    )
    exporter.add_argument("workbook", type=Path)
    exporter.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    try:
        if args.command == "build":
            print(build(args.workbook, args.summaries, args.output))
        elif args.command == "check":
            print(f"OK: {len(check(args.workbook))} decisions")
            if args.render:
                print(f"Render fit OK: {check_rendering(args.workbook)} rows")
        else:
            content = (
                json.dumps(check(args.workbook), ensure_ascii=False, indent=2) + "\n"
            )
            if args.output:
                distinct(args.workbook, args.output)
                args.output.write_text(content, encoding="utf-8")
            else:
                print(content, end="")
    except (ValueError, OSError, KeyError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
