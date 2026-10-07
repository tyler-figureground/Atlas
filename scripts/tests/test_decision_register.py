"""Temporary-file integration tests for the decision register CLI."""

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT = Path(__file__).resolve().parents[1] / "decision_register.py"
SPEC = importlib.util.spec_from_file_location("decision_register", SCRIPT)
register = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(register)


class DecisionRegisterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "register.xlsx"
        self.output = self.root / "team.xlsx"
        self.summaries = self.root / "summaries.json"
        book = Workbook()
        sheet = book.active
        sheet.title = register.DETAIL
        for col, name in enumerate(register.DETAIL_HEADERS, 1):
            sheet.cell(4, col, name)
        for row, key, status, owner in [
            (5, "D01", "Open", "Alex"),
            (7, "D02", "Decided", "Sam"),
        ]:
            sheet.cell(row, 1, key)
            sheet.cell(row, 2, "Original title")
            sheet.cell(row, 6, status)
            sheet.cell(row, 10, owner)
            sheet.cell(row, 16, "=SUM(1,2)")
            sheet.cell(row, 2).font = Font(name="Arial", size=13, bold=True)
            sheet.cell(row, 2).fill = PatternFill("solid", fgColor="FFFF00")
            sheet.cell(row, 2).alignment = Alignment(wrap_text=True)
            sheet.cell(row, 16).number_format = "$#,##0.00"
        sheet.merge_cells("A1:V1")
        sheet["A1"] = "Original register"
        sheet.column_dimensions["B"].width = 62
        sheet.row_dimensions[5].height = 80
        sheet.freeze_panes = "C5"
        book.save(self.source)
        book.close()
        self.items = [
            {
                "id": "D01",
                "topic": "Facade",
                "position": "Review options",
                "next": "Alex: Friday",
            },
            {"id": "D02", "topic": "Roof", "position": "Approved", "next": ""},
        ]
        self.write_summaries()
        self.original = self.source.read_bytes()

    def write_summaries(self):
        self.summaries.write_text(json.dumps(self.items), encoding="utf-8")

    def build(self):
        return register.build(self.source, self.summaries, self.output)

    def mutate(self, change):
        book = load_workbook(self.output)
        change(book[register.TEAM], book)
        book.save(self.output)
        book.close()

    def test_preserves_canonical_values_styles_formulas_and_layout(self):
        self.build()
        self.assertEqual(self.original, self.source.read_bytes())
        before, after = load_workbook(self.source), load_workbook(self.output)
        try:
            a, b = before[register.DETAIL], after[register.DETAIL]
            for row in a:
                for cell in row:
                    other = b[cell.coordinate]
                    self.assertEqual(cell.value, other.value)
                    self.assertEqual(cell.data_type, other.data_type)
                    for field in (
                        "font",
                        "fill",
                        "border",
                        "alignment",
                        "protection",
                        "number_format",
                    ):
                        self.assertEqual(
                            copy(getattr(cell, field)), copy(getattr(other, field))
                        )
            self.assertEqual(str(a.merged_cells), str(b.merged_cells))
            self.assertEqual(
                a.column_dimensions["B"].width, b.column_dimensions["B"].width
            )
            self.assertEqual(a.row_dimensions[5].height, b.row_dimensions[5].height)
            self.assertEqual(a.freeze_panes, b.freeze_panes)
            self.assertEqual(after.active.title, register.TEAM)
            self.assertEqual(
                after[register.TEAM]["A6"].hyperlink.location, "'Decision Register'!A7"
            )
        finally:
            before.close()
            after.close()
        self.assertEqual(register.check(self.output), self.items)

    def test_invalid_summaries_leave_input_and_output_untouched(self):
        original_items = self.items
        cases = [
            [],
            original_items + [original_items[0]],
            [dict(original_items[0], id="other"), original_items[1]],
        ]
        for field, value in [
            ("topic", "x" * 56),
            ("position", "x" * 151),
            ("next", "x" * 101),
            ("topic", " "),
            ("position", 2),
            ("next", ""),
        ]:
            cases.append([dict(original_items[0], **{field: value}), original_items[1]])
        for items in cases:
            with self.subTest(items=items):
                self.items = items
                self.write_summaries()
                with self.assertRaises(ValueError):
                    self.build()
                self.assertFalse(self.output.exists())
                self.assertEqual(self.original, self.source.read_bytes())
        self.output.write_bytes(b"existing output")
        with self.assertRaises(ValueError):
            self.build()
        self.assertEqual(self.output.read_bytes(), b"existing output")

    def test_same_path_rejected(self):
        with self.assertRaises(ValueError):
            register.build(self.source, self.summaries, self.source)
        self.assertEqual(self.original, self.source.read_bytes())

    def test_saved_value_validation(self):
        for address, value in [
            ("A5", "D02"),
            ("A5", "unknown"),
            ("D5", "Decided"),
            ("E5", "Invented"),
            ("B5", "x" * 56),
            ("C5", "x" * 151),
            ("F5", "x" * 101),
            ("F5", None),
            ("C5", "=1+1"),
        ]:
            with self.subTest(address=address, value=value):
                self.build()
                self.mutate(
                    lambda s, b, address=address, value=value: setattr(
                        s[address], "value", value
                    )
                )
                with self.assertRaises(ValueError):
                    register.check(self.output)

    def test_formatting_failures(self):
        from openpyxl.worksheet.filters import FilterColumn, Filters

        changes = {
            "width": lambda s, b: setattr(s.column_dimensions["C"], "width", 50),
            "height zero": lambda s, b: setattr(s.row_dimensions[5], "height", 0),
            "height tall": lambda s, b: setattr(s.row_dimensions[5], "height", 49),
            "height short": lambda s, b: setattr(s.row_dimensions[5], "height", 35),
            "hidden row": lambda s, b: setattr(s.row_dimensions[5], "hidden", True),
            "hidden column": lambda s, b: setattr(
                s.column_dimensions["G"], "hidden", True
            ),
            "wrap": lambda s, b: setattr(
                s["C5"], "alignment", Alignment(vertical="top")
            ),
            "font": lambda s, b: setattr(s["C5"], "font", Font(size=10)),
            "zoom": lambda s, b: setattr(s.sheet_view, "zoomScale", 70),
            "freeze": lambda s, b: setattr(s, "freeze_panes", "A4"),
            "active": lambda s, b: setattr(b, "active", 1),
            "filter range": lambda s, b: setattr(s.auto_filter, "ref", "A4:F5"),
            "filter criteria": lambda s, b: s.auto_filter.filterColumn.append(
                FilterColumn(colId=3, filters=Filters(filter=["Open"]))
            ),
            "filter mode": lambda s, b: setattr(s.sheet_properties, "filterMode", True),
            "group": lambda s, b: setattr(s.row_dimensions[5], "outlineLevel", 1),
            "indicators": lambda s, b: setattr(
                s.sheet_view, "showOutlineSymbols", True
            ),
            "link": lambda s, b: setattr(s["A5"], "hyperlink", None),
            "merge": lambda s, b: s.merge_cells("B5:C5"),
        }
        for name, change in changes.items():
            with self.subTest(name=name):
                self.build()
                self.mutate(change)
                with self.assertRaises(ValueError):
                    register.check(self.output)

    def test_unsafe_existing_view_rejected(self):
        for change in [
            lambda s, b: s.merge_cells("A1:F1"),
            lambda s, b: setattr(s["C5"], "value", "=1+1"),
            lambda s, b: setattr(s["G5"], "value", "extra"),
            lambda s, b: setattr(s["A1"], "value", "Other view"),
        ]:
            self.build()
            self.mutate(change)
            prior = self.output.read_bytes()
            target = self.root / "rebuilt.xlsx"
            with self.assertRaises(ValueError):
                register.build(self.output, self.summaries, target)
            self.assertFalse(target.exists())
            self.assertEqual(prior, self.output.read_bytes())

    def test_canonical_validation(self):
        for address, value in [
            ("A7", "D01"),
            ("A7", None),
            ("J5", "x" * 41),
            ("F5", "=1"),
            ("A4", "Wrong"),
        ]:
            with self.subTest(address=address):
                self.source.write_bytes(self.original)
                book = load_workbook(self.source)
                book[register.DETAIL][address] = value
                book.save(self.source)
                book.close()
                with self.assertRaises(ValueError):
                    self.build()
                self.assertFalse(self.output.exists())

    def test_rebuild_after_canonical_changes_and_new_record(self):
        self.build()

        def update(s, book):
            detail = book[register.DETAIL]
            detail["F5"] = "Decided"
            detail["J5"] = "New owner"
            for col, value in ((1, "D03"), (2, "New topic"), (6, "Open")):
                detail.cell(8, col, value)

        self.mutate(update)
        self.items.append(
            {
                "id": "D03",
                "topic": "New topic",
                "position": "Under review",
                "next": "Confirm choice",
            }
        )
        self.write_summaries()
        rebuilt = self.root / "rebuilt.xlsx"
        register.build(self.output, self.summaries, rebuilt)
        self.assertEqual(len(register.check(rebuilt)), 3)
        book = load_workbook(rebuilt)
        self.assertEqual(book[register.TEAM]["D5"].value, "Decided")
        self.assertEqual(book[register.TEAM]["E5"].value, "New owner")
        book.close()

    def test_rebuild_preserves_comments_on_populated_and_blank_cells(self):
        from openpyxl.comments import Comment

        self.build()

        def annotate(sheet, book):
            for address in ("C5", "F6", "A1"):
                sheet[address].comment = Comment("Keep this human note", "Reviewer")

        self.mutate(annotate)
        rebuilt = self.root / "rebuilt.xlsx"
        register.build(self.output, self.summaries, rebuilt)
        book = load_workbook(rebuilt)
        for address in ("C5", "F6", "A1"):
            self.assertEqual(
                book[register.TEAM][address].comment.text, "Keep this human note"
            )
        book.close()
        self.mutate(
            lambda sheet, book: setattr(
                sheet["C9"], "comment", Comment("Unkeyed note", "Reviewer")
            )
        )
        with self.assertRaises(ValueError):
            register.build(self.output, self.summaries, self.root / "unsafe.xlsx")

    def test_blank_owner_stays_unknown(self):
        book = load_workbook(self.source)
        book[register.DETAIL]["J5"] = None
        book.save(self.source)
        book.close()
        self.build()
        self.assertEqual(register.check(self.output), self.items)
        saved = load_workbook(self.output)
        self.assertIsNone(saved[register.TEAM]["E5"].value)
        saved.close()

    @unittest.skipUnless(
        importlib.util.find_spec("aspose"),
        "optional spreadsheet renderer not installed",
    )
    def test_render_measurement_catches_clipped_text(self):
        self.build()
        self.assertEqual(register.check_rendering(self.output), 2)
        self.mutate(lambda s, b: setattr(s["F5"], "value", "W" * 100))
        register.check(
            self.output
        )  # Within the character budget, but not the line budget.
        with self.assertRaisesRegex(ValueError, "would clip"):
            register.check_rendering(self.output)

    def test_markdown_export_keeps_all_history(self):
        sys.path.insert(0, str(SCRIPT.parent))
        self.addCleanup(lambda: sys.path.remove(str(SCRIPT.parent)))
        from export_decision_register import export

        history = (
            "[2026-01-01] " + "Full evidence " * 900 + "\n[2026-02-01] Later evidence"
        )
        book = load_workbook(self.source)
        book[register.DETAIL]["V5"] = history
        book[register.DETAIL]["A5"] = 0
        self.items[0]["id"] = 0
        self.write_summaries()
        book.save(self.source)
        book.close()
        self.build()
        summary = self.root / "REGISTER.md"
        self.assertEqual(export(self.output, summary), 2)
        self.assertIn(history, (self.root / "DETAILS.md").read_text(encoding="utf-8"))
        self.assertNotIn(history, summary.read_text(encoding="utf-8"))
        self.assertIn("[0](DETAILS.md#0)", summary.read_text(encoding="utf-8"))
        with self.assertRaises(ValueError):
            export(self.output, self.root / "DETAILS.md")

    def test_cli_roundtrip_and_literal_formula_like_text(self):
        self.items[0]["position"] = "=Literal summary, not a formula"
        self.write_summaries()
        result = subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                "build",
                str(self.source),
                "--summaries",
                str(self.summaries),
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        built = self.source.with_name("register.team.xlsx")
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "export-summaries", str(built)],
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), self.items)
        self.summaries.write_text(result.stdout, encoding="utf-8")
        register.build(built, self.summaries, self.output)
        self.assertEqual(register.check(self.output), self.items)
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "check", str(self.output)],
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
